from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SKY = "0284C7"
WASH = "E6F7FF"
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xB4, 0x53, 0x09)


def shade_cell(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run(run, text, size=11, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_text(paragraph, text, size=11, bold=False, color=INK, font="Calibri"):
    set_run(paragraph.add_run(), text, size, bold, color, font)
    return paragraph


def no_border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tblPr.append(borders)


def heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    add_text(p, f"{number}   {title}", 14, True, RGBColor(0x02, 0x84, 0xC7))
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text, 11)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "•  " + text, 11)
    return p


def line(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, label, 11, True)
    add_text(p, "  ________________________________", 11, False, MUTED)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.4)
section.bottom_margin = Cm(1.4)
section.left_margin = Cm(1.6)
section.right_margin = Cm(1.6)

banner = doc.add_table(rows=1, cols=1)
banner.autofit = True
cell = banner.cell(0, 0)
shade_cell(cell, SKY)
cell.text = ""
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(0)
add_text(p, "ZEYHAAN", 28, True, WHITE)
p2 = cell.add_paragraph()
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after = Pt(2)
add_text(p2, "Mkataba wa utoaji wa PhoneVault", 14, False, WHITE)
p3 = cell.add_paragraph()
p3.paragraph_format.space_before = Pt(0)
p3.paragraph_format.space_after = Pt(8)
add_text(p3, "Till ya duka la simu  ·  kaunta moja  ·  16 Oktoba 2026", 10, False, RGBColor(0xE6, 0xF7, 0xFF))

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(12)
intro.paragraph_format.space_after = Pt(8)
add_text(intro, "Mkataba huu unaeleza mfumo unaokabidhiwa, bei yake, na ada ya seva ambayo mteja anailipa mwenyewe kila mwaka.", 11)

price = doc.add_table(rows=2, cols=2)
price.autofit = True
left, right = price.rows[0].cells
shade_cell(left, WASH)
shade_cell(right, "FFF7ED")
left.text = ""
right.text = ""
a = left.paragraphs[0]
add_text(a, "BEI YA MFUMO", 9, True, MUTED)
b = left.add_paragraph()
add_text(b, "TSH 400,000", 22, True, RGBColor(0x02, 0x84, 0xC7))
c = left.add_paragraph()
add_text(c, "Laki nne, malipo ya mara moja", 10, False, MUTED)
d = right.paragraphs[0]
add_text(d, "SEVA KWA MWAKA", 9, True, AMBER)
e = right.add_paragraph()
add_text(e, "Mteja analipa", 18, True, AMBER)
f = right.add_paragraph()
add_text(f, "Zeyhaan hailipi ada ya seva", 10, False, MUTED)

heading(doc, "01", "Wahusika")
body(doc, "Mtoa huduma ni Zeyhaan. Mteja ni mmiliki wa duka la simu anayesaini chini ya mkataba huu. Mteja anajaza taarifa zake kabla ya kusaini.")
line(doc, "Jina la mteja")
line(doc, "Jina la duka")
line(doc, "Simu")
line(doc, "Mahali")

heading(doc, "02", "Bidhaa inayokabidhiwa")
body(doc, "Zeyhaan inamkabidhi mteja PhoneVault, programu ya till inayotumika kwenye simu ya kaunta moja. Mauzo, stoo, deni, risiti, marejesho na uchambuzi vinakaa kwenye duka. Programu haikatii pesa za mteja. Mteja wa duka hulipa cash, au hulipa Lipa Namba ya duka nje ya programu.")

heading(doc, "03", "Kazi za mfumo")
for item in [
    "Kuuza simu na vifaa, ikiwemo IMEI, vikundi na stoo.",
    "Malipo ya cash na Lipa Namba, deni zima na deni la sehemu.",
    "Kikokotoo cha malipo: kinaandika alichotoa mteja na chenji aliyorudishiwa, kwenye skrini na kwenye risiti. Chenji haihesabiwi kuwa mauzo.",
    "Risiti ya PDF inayofungua WhatsApp kwa namba iliyo kwenye mauzo. Muuzaji anabonyeza Tuma ndani ya WhatsApp.",
    "Kurudisha bidhaa kwa sababu. Stoo, deni, faida na ripoti zinajisasisha, na taarifa inaandikwa.",
    "Uchambuzi wa siku 7, wiki iliyopita, mwezi, mwezi uliopita, mwaka, mwaka uliopita, na tarehe yoyote. Unaona gharama na faida ya kila bidhaa.",
    "Mmiliki na muuzaji, kila mmoja na PIN yake, na sehemu anazoruhusiwa.",
    "Mauzo yanaendelea hata mtandao ukikatika. Nakala ya duka inaenda kwenye kompyuta ya duka mtandao ukirudi.",
    "Mipangilio ya jina la duka, mahali, simu ya duka, ujumbe wa chini ya risiti, na Lipa Namba.",
]:
    bullet(doc, item)

heading(doc, "04", "Bei")
body(doc, "Bei ya PhoneVault ni shilingi laki nne, yaani TSH 400,000. Ni malipo ya mara moja, yanayolipwa na mteja kwa Zeyhaan, kwa kukabidhiwa mfumo huu na kuonyeshwa namna ya kuutumia kwenye kaunta moja.")
body(doc, "Bei hii ndiyo ada ya bidhaa. Haifichi ada ya seva, ada ya SMS, wala gharama ya mtandao.")

heading(doc, "05", "Ada ya seva kwa mwaka")
body(doc, "Seva inayohifadhi nakala ya duka, na ada yake ya kila mwaka, inalipwa na mteja mwenyewe. Zeyhaan hailipi ada hiyo, haichukui deni la ada hiyo, na haijumuishi ada hiyo ndani ya TSH 400,000.")
body(doc, "Mteja anachagua na kulipa mtoa seva moja kwa moja. Mwaka mpya ukifika, mteja anahakikisha ada imelipwa ili nakala iendelee. Zeyhaan haina wajibu wa kulipa bili ya seva kwa niaba ya mteja.")

heading(doc, "06", "Gharama nyingine za mteja")
bullet(doc, "Lipa Namba ni namba ya duka. Pesa zinaingia kwa mteja, si kwa Zeyhaan.")
bullet(doc, "Mtandao wa simu na kompyuta ya duka ni wa mteja.")
bullet(doc, "Endapo mteja ataweka akaunti ya SMS, ada ya ujumbe inalipwa na mteja kwa kampuni ya SMS.")

heading(doc, "07", "Yasiyo ndani ya mkataba huu")
bullet(doc, "Risiti ya kodi ya TRA.")
bullet(doc, "Printer ya joto.")
bullet(doc, "Kaunta ya pili, tawi lingine, au simu ya pili inayouza pamoja na ya kwanza.")
body(doc, "Kitu chochote kati ya hivyo, kikibidi baadaye, kinakuwa maagizo mapya na bei yake mpya.")

heading(doc, "08", "Ukabidhi")
body(doc, "Zeyhaan inamkabidhi mteja programu iliyowekwa kwenye simu ya kaunta, na kumwonyesha kuuza, stoo, deni, risiti, marejesho na uchambuzi. Mteja anaanza kuitumia baada ya kulipa TSH 400,000.")
body(doc, "Mteja anahifadhi PIN ya mmiliki, na anabadilisha PIN ya majaribio kabla ya kuanza kazi. Mteja anaweka jina halisi la duka, mahali, simu na Lipa Namba kwenye mipangilio.")

heading(doc, "09", "Wajibu wa kila upande")
body(doc, "Zeyhaan inakabidhi mfumo unaofanya kazi zilizoandikwa kifungu cha 03, na anamwelekeza mteja siku ya ukabidhi. Mteja anatumia mfumo kwenye duka lake, anaingiza bidhaa zake mwenyewe, na analipa ada ya seva kila mwaka.")
body(doc, "Taarifa za mauzo, wateja na stoo ni mali ya mteja. Zeyhaan haitumii taarifa hizo kwa biashara nyingine.")

heading(doc, "10", "Sahihi")
body(doc, "Kwa kusaini, pande zote mbili zinakubali mkataba huu. Nakala mbili, kila upande unabaki na moja.")

sign = doc.add_table(rows=1, cols=2)
sign.autofit = True
for cell, title in zip(sign.rows[0].cells, ("Zeyhaan", "Mteja")):
    shade_cell(cell, WASH)
    cell.text = ""
    h = cell.paragraphs[0]
    add_text(h, title, 13, True, RGBColor(0x02, 0x84, 0xC7))
    for label in ("Jina", "Sahihi", "Tarehe"):
        q = cell.add_paragraph()
        q.paragraph_format.space_before = Pt(10)
        add_text(q, label, 10, True, MUTED)
        r = cell.add_paragraph()
        add_text(r, "________________________", 11, False, MUTED)

foot = doc.add_paragraph()
foot.paragraph_format.space_before = Pt(16)
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(foot, "Zeyhaan  ·  PhoneVault  ·  TSH 400,000  ·  seva ya mwaka inalipwa na mteja", 9, False, MUTED)

out = r"C:\Users\USER\.gemini\antigravity-ide\scratch\phone-shop-system\Zeyhaan.docx"
doc.save(out)
print(out)
